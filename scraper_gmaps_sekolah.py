from playwright.sync_api import sync_playwright
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
import time
import urllib.parse
import re
import os
import json
import sys
import signal
from datetime import datetime

CHECKPOINT_FILE = "scraper_gmaps_checkpoint.json"
DATABASE_JSON_FILE = "database_sekolah.json"
GLOBAL_LEADS = []
GLOBAL_KEYWORD = "sekolah"

def save_to_professional_excel(data, keyword):
    if not data:
        print("Tidak ada data untuk diekspor ke Excel.")
        return
        
    df = pd.DataFrame(data)
    df.drop_duplicates(subset=['Nama Instansi', 'Alamat Lengkap'], keep='first', inplace=True)
    df.insert(0, 'No', range(1, len(df) + 1))
    
    timestamp_str = datetime.now().strftime("%d_%m_%Y_%H_%M_%S")
    output_name = f"Scraper_Gmaps_Sekolah_{keyword.capitalize()}_{timestamp_str}.xlsx"
    
    with pd.ExcelWriter(output_name, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Database Leads')
        
    wb = openpyxl.load_workbook(output_name)
    ws = wb['Database Leads']
    
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    
    header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    zebra_fill = PatternFill(start_color="F8F9F9", end_color="F8F9F9", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='D5D8DC'), right=Side(style='thin', color='D5D8DC'),
        top=Side(style='thin', color='D5D8DC'), bottom=Side(style='thin', color='D5D8DC')
    )
    
    for col_num in range(1, len(df.columns) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
    
    ws.row_dimensions[1].height = 28
    
    for row_num in range(2, len(df) + 2):
        ws.row_dimensions[row_num].height = 20
        current_fill = zebra_fill if (row_num % 2 == 0) else white_fill
        
        for col_num in range(1, len(df.columns) + 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.fill = current_fill
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)
            
            col_name = df.columns[col_num - 1]
            if col_name in ["No", "NPSN", "Kota", "Status"]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')
                
    status_col_letter = get_column_letter(df.columns.get_loc('Status') + 1)
    dv = DataValidation(type="list", formula1='"New, Contacted, Follow Up, Deal, Rejected"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{status_col_letter}2:{status_col_letter}{len(df) + 1}")
    
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row == 1:
                max_len = max(max_len, len(str(cell.value or '')))
            else:
                val_str = str(cell.value or '')
                if len(val_str) > 40:
                    max_len = max(max_len, 40)
                else:
                    max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)
        
    wb.save(output_name)
    print(f"Selesai! {len(df)} data unik berhasil diekspor ke file Excel: '{output_name}'")

def signal_handler(sig, frame):
    print("\nProgram dihentikan paksa (Ctrl + C). Mengekspor data yang ada ke Excel...")
    save_to_professional_excel(GLOBAL_LEADS, GLOBAL_KEYWORD)
    sys.exit(0)

signal.signal(signal.SIGINT, signal_handler)

def load_database():
    if os.path.exists(DATABASE_JSON_FILE):
        try:
            with open(DATABASE_JSON_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
                print(f"Ditemukan database JSON: {len(data.get('leads', []))} data sekolah sudah tersimpan.")
                return data
        except Exception:
            pass
    return {"leads": [], "processed_names": []}

def save_to_database(leads, processed_names):
    try:
        with open(DATABASE_JSON_FILE, "w", encoding="utf-8") as f:
            json.dump({"leads": leads, "processed_names": processed_names}, f, ensure_ascii=False, indent=4)
    except Exception as e:
        print(f"Gagal menyimpan ke JSON: {e}")

def scrape_gmaps_sekolah(keyword, max_results=200):
    global GLOBAL_LEADS, GLOBAL_KEYWORD
    GLOBAL_KEYWORD = keyword
    
    db_data = load_database()
    leads = db_data["leads"]
    GLOBAL_LEADS = leads
    processed_names = set(db_data["processed_names"])
    
    print(f"Memulai ekstraksi Google Maps untuk wilayah: 'Sekolah di {keyword}'")
    
    with sync_playwright() as p:
        user_data_dir = os.path.expanduser("~") + "/playwright_chrome_profile"
        
        print(f"Membuka browser Google Chrome dengan profil di: {user_data_dir}")
        context = p.chromium.launch_persistent_context(
            user_data_dir=user_data_dir,
            headless=False,
            channel="chrome",
            ignore_https_errors=True,
            args=["--start-maximized"]
        )
        
        page = context.new_page()
        
        search_query = f"sekolah di {keyword}"
        encoded_query = urllib.parse.quote(search_query)
        gmaps_url = f"https://www.google.com/maps/search/{encoded_query}"
        
        try:
            print(f"Mengakses: {gmaps_url}")
            page.goto(gmaps_url, timeout=60000)
            
            feed_selector = 'div[role="feed"]'
            try:
                page.wait_for_selector(feed_selector, timeout=20000)
            except:
                print("Panel feed Google Maps tidak ditemukan.")
            
            time.sleep(4)
            
            print("Menggulir panel Google Maps untuk memuat data...")
            scroll_attempts = 0
            max_scroll_without_change = 5
            last_height_count = 0
            
            while len(leads) < max_results and scroll_attempts < max_scroll_without_change:
                try:
                    page.evaluate("""
                        const feed = document.querySelector('div[role="feed"]');
                        if (feed) {
                            feed.scrollBy(0, 1000);
                        }
                    """)
                    time.sleep(2.5)
                    
                    current_listings = page.locator('div[role="feed"] a.hfpxzc').all()
                    current_count = len(current_listings)
                    
                    if current_count == last_height_count:
                        scroll_attempts += 1
                    else:
                        scroll_attempts = 0
                        last_height_count = current_count
                        
                    if "You've reached the end of the list" in page.evaluate("() => document.body.innerText"):
                        print("Sudah mencapai akhir daftar hasil pencarian Google Maps.")
                        break
                        
                except Exception as e:
                    print(f"Kesalahan saat scrolling: {e}")
                    break
            
            listings = page.locator('div[role="feed"] a.hfpxzc').all()
            print(f"Ditemukan {len(listings)} tautan di Google Maps. Memulai ekstraksi detail...")
            
            for i, listing in enumerate(listings):
                if len(leads) >= max_results:
                    break
                try:
                    
                    listing.scroll_into_view_if_needed()
                    nama_mentah = listing.get_attribute("aria-label") or f"Sekolah {i+1}"
                    
                    # Bersihkan teks ".Visited link" dari nama sekolah
                    nama_instansi = re.sub(r'\s*·?\s*Visited link\s*', '', nama_mentah, flags=re.IGNORECASE).strip()
                    
                    nama_lower = nama_instansi.lower()
                    valid_pattern = r'\b(paud|tk|tka|tkb|kb|playgroup|sd|sdn|sdi|smp|smpn|smpi|mts|mtsn|sma|sman|smai|smk|smkn|ma|man|mak|min|slb|universitas|institut|politeknik|akademi|sekolah|pesantren|ponpes)\b'
                    invalid_pattern = r'\b(lpk|kursus|les|bimbel|toko|agen|pt|cv|yayasan|pelatihan|bengkel|tour|travel|printing|fotocopy|jasa|sewa|warung|koperasi|klinik|apotek|salon)\b'
                    
                    if not re.search(valid_pattern, nama_lower) or re.search(invalid_pattern, nama_lower):
                        print(f"[{i+1}/{len(listings)}] Melewati: {nama_instansi} (Bukan institusi pendidikan target)")
                        continue
                    
                    if nama_instansi in processed_names:
                        print(f"[{i+1}/{len(listings)}] Melewati: {nama_instansi} (Sudah ada di database JSON)")
                        continue
                        
                    print(f"[{i+1}/{len(listings)}] Mengekstrak: {nama_instansi}")
                    listing.click(timeout=5000)
                    time.sleep(3)
                    
                    try:
                        page.keyboard.press("Tab") 
                        time.sleep(0.5)
                        for _ in range(8):
                            page.keyboard.press("PageDown")
                            time.sleep(1.2)
                    except:
                        pass
                    
                    time.sleep(3)
                    gmaps_link = page.url
                    
                    website = "-"
                    instagram = "-"
                    kemendikdasmen_link = "-"
                    
                    try:
                        website_elem = page.locator('a[data-item-id="authority"]')
                        if website_elem.count() > 0:
                            website = website_elem.first.get_attribute("href") or "-"
                    except:
                        pass
                        
                    html_content = ""
                    try:
                        html_content += page.content() + "\n"
                    except:
                        pass
                        
                    for frame in page.frames:
                        try:
                            html_content += frame.content() + "\n"
                        except:
                            pass
                    
                    unquoted_html = urllib.parse.unquote(html_content)
                    
                    if kemendikdasmen_link == "-":
                        # Regex diperbarui agar wajib menangkap path /pendidikan/... di belakangnya
                        kem_match = re.search(r'(https?://[a-zA-Z0-9.-]*kemendikdasmen\.go\.id/[^\s"\'<>\\&]+)', unquoted_html)
                        if kem_match:
                            kemendikdasmen_link = kem_match.group(1).rstrip('.,;)')
                            
                    if instagram == "-":
                        ig_match = re.search(r'(https?://(?:www\.)?instagram\.com/[A-Za-z0-9_.]+)', unquoted_html)
                        if ig_match:
                            instagram = ig_match.group(1).rstrip('.')
                            
                    if website == "-":
                        web_matches = re.findall(r'(https?://[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}[^\s"\'<>\\\\]*)', unquoted_html)
                        # Ditambahkan 'schema.org' ke exclusions agar tidak masuk ke kolom website
                        exclusions = ["google", "gstatic", "facebook", "youtube", "wikipedia.org", "kemendikbud", "kemendikdasmen", "instagram", "twitter", "tiktok", "linkedin", "w3.org", "playwright", "schema.org"]
                        for w in web_matches:
                            clean_w = w.split('&')[0].rstrip('.,;)')
                            if not any(excl in clean_w.lower() for excl in exclusions):
                                website = clean_w
                                break

                    detail_text = ""
                    try:
                        detail_text = page.locator('div[role="main"]').inner_text(timeout=5000)
                    except:
                        detail_text = page.evaluate("() => document.body.innerText")
                        
                    for frame in page.frames:
                        try:
                            detail_text += " " + frame.locator('body').inner_text()
                        except:
                            pass
                    
                    npsn_match = re.search(r'(?i)NPSN[\s\:\-]*(\d{8})', detail_text)
                    if not npsn_match:
                        npsn_match = re.search(r'\b(\d{8})\b', detail_text)
                    npsn = npsn_match.group(1) if npsn_match else "-"
                    
                    # EKSTRAKSI ALAMAT DARI OVERVIEW ADDRESS (Elemen spesifik Google Maps)
                    alamat = "-"
                    try:
                        alamat_elem = page.locator('button[data-item-id="address"] div.Io6YTe')
                        if alamat_elem.count() > 0:
                            alamat = alamat_elem.first.inner_text().strip()
                    except:
                        pass
                    
                    # Fallback alamat jika elemen spesifik gagal terambil
                    if alamat == "-":
                        alamat_match = re.search(r'([A-Za-z0-9\.\,\s\-]+(?:Jl\.|Jalan|Kp\.|Kec\.|Kab\.|Kota|Prov\.)[^O\n]+)', detail_text)
                        alamat = alamat_match.group(1).strip() if alamat_match else f"Wilayah {keyword.capitalize()}"
                        alamat = re.sub(r'Open(\s*now)?\s*$', '', alamat, flags=re.IGNORECASE).strip()
                    
                    # EKSTRAKSI TELEPON DARI OVERVIEW (Elemen spesifik Google Maps)
                    telp = "-"
                    try:
                        telp_elem = page.locator('button[data-item-id^="phone:tel:"] div.Io6YTe')
                        if telp_elem.count() > 0:
                            telp_raw = telp_elem.first.inner_text().strip()
                            telp = re.sub(r'[^\d+]', '', telp_raw)
                    except:
                        pass
                    
                    # Fallback regex telepon
                    if telp == "-":
                        telp_match = re.search(r'(\+?62|0)[0-9\-\s]{8,13}', detail_text)
                        telp = telp_match.group(0).strip() if telp_match else "-"
                        if telp != "-" and telp.startswith("0"):
                            telp = "62" + re.sub(r'[\s\-]', '', telp[1:])
                        elif telp.startswith("+62"):
                            telp = "62" + re.sub(r'[\s\-]', '', telp[3:])
                        
                    email_match = re.search(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}', detail_text)
                    email = email_match.group(0) if email_match else "-"
                    
                    if email == "-":
                        links_to_check = []
                        if website != "-": links_to_check.append(website)
                        if kemendikdasmen_link != "-": links_to_check.append(kemendikdasmen_link)
                        if instagram != "-": links_to_check.append(instagram)
                            
                        for target_link in links_to_check:
                            if email != "-": break
                            print(f"Mencari email di tab baru: {target_link}")
                            try:
                                temp_page = context.new_page()
                                temp_page.goto(target_link, timeout=12000, wait_until="domcontentloaded")
                                time.sleep(2)
                                temp_text = temp_page.evaluate("() => document.body.innerText")
                                e_match = re.search(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}', temp_text)
                                if e_match:
                                    email = e_match.group(0)
                                    print(f"Berhasil menemukan email: {email}")
                            except Exception:
                                pass
                            finally:
                                try:
                                    temp_page.close()
                                except: pass
                    
                    new_lead = {
                        "Nama Instansi": nama_instansi,
                        "NPSN": npsn,
                        "Bidang Usaha / Kategori": "Sekolah",
                        "No. WA / Telephon": telp,
                        "Alamat Email": email,
                        "Alamat Lengkap": alamat,
                        "Kota": keyword.capitalize(),
                        "Link Gmaps": gmaps_link,
                        "Link Website": website,
                        "Akun Media Sosial": instagram,
                        "Link Kemendikdasmen": kemendikdasmen_link,
                        "Sumber": "Google Maps",
                        "Status": "New"
                    }
                    
                    leads.append(new_lead)
                    GLOBAL_LEADS = leads
                    processed_names.add(nama_instansi)
                    
                    save_to_database(leads, list(processed_names))
                    
                except Exception as e:
                    print(f"Gagal mengekstrak item ke-{i+1}: {e}")
                    continue
                    
        except Exception as e:
            print(f"Error pada browser: {e}")
        finally:
            context.close()
            
    return leads

if __name__ == "__main__":
    print("="*50)
    print("PROGRAM SCRAPER GOOGLE MAPS SEKOLAH")
    print("="*50)
    print("Pilih Wilayah Target Pencarian:")
    print("1. Salatiga")
    print("2. Boyolali")
    print("3. Solo")
    print("4. Semarang")
    print("5. Kabupaten Semarang")
    print("="*50)
    
    pilihan = input("Masukkan pilihan angka (1-5): ").strip()
    
    kota_dict = {
        "1": "salatiga",
        "2": "boyolali",
        "3": "solo",
        "4": "semarang",
        "5": "kabupaten semarang"
    }
    
    if pilihan in kota_dict:
        target_kota = kota_dict[pilihan]
        print(f"Anda memilih: {target_kota.capitalize()}")
        
        hasil = scrape_gmaps_sekolah(target_kota, max_results=100)
        save_to_professional_excel(hasil, target_kota)
    else:
        print("Pilihan tidak valid. Silakan jalankan ulang program.")